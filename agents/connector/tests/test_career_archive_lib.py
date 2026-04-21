"""Pure helpers for career-archive-import.py — Stage 1: walk the career
archive roots, extract text per file type, batch for LLM classification,
merge classified records into the archive-index.json idempotently.

Tests cover the library surface; the runner lives in
scripts/career-archive-import.py and is exercised via integration run.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

_SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(_SCRIPTS_DIR))

from career_archive_lib import (  # type: ignore
    ALLOWED_CLASSES,
    ALLOWED_EXTENSIONS,
    SKIP_EXTENSIONS,
    build_classification_prompt,
    extract_text,
    infer_role_from_path,
    iter_file_records,
    merge_records_into_index,
    parse_classification_response,
)


# --- iter_file_records ---

# Helper: most tests need files >= MIN_SIZE_BYTES (200) — the walker filters
# tiny files as low-signal. Pad with plausible bio-ish content.
_PAD = ("Sam Smith led marketplace data science. Launched pricing guidance. " * 10)


def test_iter_file_records_yields_files_with_role_ext_mtime(tmp_path: Path):
    role_root = tmp_path / "Example Corp"
    role_root.mkdir()
    (role_root / "Roadmap.docx").write_text(_PAD, encoding="utf-8")
    (role_root / "Notes.txt").write_text(_PAD, encoding="utf-8")

    records = list(iter_file_records(role_root, role="airbnb"))
    names = {r["path"].name for r in records}
    assert names == {"Roadmap.docx", "Notes.txt"}
    for r in records:
        assert r["role"] == "airbnb"
        assert r["ext"] in {".docx", ".txt"}
        assert r["size"] > 0
        assert "T" in r["mtime"]  # ISO 8601


def test_iter_file_records_recurses_into_subdirs(tmp_path: Path):
    (tmp_path / "Exit").mkdir()
    (tmp_path / "Exit" / "Departure.pdf").write_text(_PAD, encoding="utf-8")
    (tmp_path / "Top.docx").write_text(_PAD, encoding="utf-8")
    paths = {r["path"].name for r in iter_file_records(tmp_path, role="airbnb")}
    assert paths == {"Departure.pdf", "Top.docx"}


def test_iter_file_records_skips_non_allowed_extensions(tmp_path: Path):
    (tmp_path / "team.jpg").write_text(_PAD, encoding="utf-8")
    (tmp_path / "recording.mov").write_text(_PAD, encoding="utf-8")
    (tmp_path / "config.xml").write_text(_PAD, encoding="utf-8")
    (tmp_path / "style.css").write_text(_PAD, encoding="utf-8")
    (tmp_path / "doc.docx").write_text(_PAD, encoding="utf-8")
    records = list(iter_file_records(tmp_path, role="airbnb"))
    names = {r["path"].name for r in records}
    assert names == {"doc.docx"}


def test_iter_file_records_missing_root_yields_nothing(tmp_path: Path):
    records = list(iter_file_records(tmp_path / "nope", role="airbnb"))
    assert records == []


def test_iter_file_records_skips_temp_office_files(tmp_path: Path):
    # Microsoft Office leaves ~$foo.docx lock files; these have no real content.
    (tmp_path / "~$Roadmap.docx").write_text(_PAD, encoding="utf-8")
    (tmp_path / "Roadmap.docx").write_text(_PAD, encoding="utf-8")
    names = {r["path"].name for r in iter_file_records(tmp_path, role="airbnb")}
    assert names == {"Roadmap.docx"}


def test_iter_file_records_skips_small_files(tmp_path: Path):
    """Tiny stub files (empty notes, name shortcuts) are low-signal;
    walker filters by MIN_SIZE_BYTES."""
    (tmp_path / "stub.txt").write_text("hi", encoding="utf-8")
    (tmp_path / "real.txt").write_text(_PAD, encoding="utf-8")
    names = {r["path"].name for r in iter_file_records(tmp_path, role="airbnb")}
    assert names == {"real.txt"}


def test_iter_file_records_skips_node_modules(tmp_path: Path):
    """Interview Prep folders sometimes contain full code checkouts.
    Walking into node_modules would explode record counts with zero signal."""
    (tmp_path / "Prep.docx").write_text(_PAD, encoding="utf-8")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "package.json").write_text("{}" * 100, encoding="utf-8")
    (tmp_path / "node_modules" / "dep").mkdir()
    (tmp_path / "node_modules" / "dep" / "index.js").write_text(_PAD, encoding="utf-8")
    names = {r["path"].name for r in iter_file_records(tmp_path, role="job-search")}
    assert names == {"Prep.docx"}


def test_iter_file_records_skips_git_dir(tmp_path: Path):
    (tmp_path / "README.md").write_text(_PAD, encoding="utf-8")
    (tmp_path / ".git").mkdir()
    (tmp_path / ".git" / "HEAD").write_text(_PAD, encoding="utf-8")
    names = {r["path"].name for r in iter_file_records(tmp_path, role="airbnb")}
    assert names == {"README.md"}


def test_iter_file_records_skips_nested_code_checkout(tmp_path: Path):
    """The actual case the operator hit: Interview Prep/Uber/Product Jam/ is a
    real code project with node_modules several levels deep."""
    deep = tmp_path / "Uber (2026)" / "Product Jam" / "supply-sim"
    deep.mkdir(parents=True)
    (deep / "README.md").write_text(_PAD, encoding="utf-8")
    nm = deep / "node_modules" / "react"
    nm.mkdir(parents=True)
    (nm / "package.json").write_text(_PAD, encoding="utf-8")
    names = {r["path"].name for r in iter_file_records(tmp_path, role="job-search")}
    assert names == {"README.md"}


# --- infer_role_from_path ---

def test_infer_role_from_path_archive_roots():
    assert infer_role_from_path(Path("E:/Dropbox/Archive/Example Corp/foo.docx")) == "airbnb"
    assert infer_role_from_path(Path("E:/Dropbox/Archive/Amazon/Exit/x.pdf")) == "amazon"
    assert infer_role_from_path(Path("E:/Dropbox/Archive/LinkedIn/y.pptx")) == "linkedin"
    assert infer_role_from_path(Path("E:/Dropbox/Archive/Twitch/Projects/z.docx")) == "twitch"


def test_infer_role_from_path_personal_roots():
    assert infer_role_from_path(Path("E:/Dropbox/Personal/Bio/LinkedIn Bio.txt")) == "bio"
    assert infer_role_from_path(Path("E:/Dropbox/Personal/Job Search/Company List.xlsx")) == "job-search"


def test_infer_role_from_path_unknown_returns_unknown():
    assert infer_role_from_path(Path("C:/random/file.docx")) == "unknown"


# --- extract_text ---

def test_extract_text_txt_utf8(tmp_path: Path):
    p = tmp_path / "bio.txt"
    p.write_text("the operator — Head of Analytics at Twitch.\nRole started 2018.", encoding="utf-8")
    text, used = extract_text(p)
    assert "the operator" in text
    assert "Twitch" in text
    assert used == "txt"


def test_extract_text_docx(tmp_path: Path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Vision: improve host outcomes via pricing.")
    doc.add_paragraph("Mission: ship pricing guidance system.")
    p = tmp_path / "roadmap.docx"
    doc.save(p)
    text, used = extract_text(p)
    assert "Vision" in text
    assert "host outcomes" in text
    assert "Mission" in text
    assert used == "docx"


def test_extract_text_xlsx(tmp_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Company", "Role", "Tier"])
    ws.append(["Stripe", "Director DS", "A"])
    ws.append(["Wayfair", "VP Data", "B"])
    p = tmp_path / "companies.xlsx"
    wb.save(p)
    text, used = extract_text(p)
    assert "Stripe" in text
    assert "Director DS" in text
    assert "Wayfair" in text
    assert used == "xlsx"


def test_extract_text_pptx(tmp_path: Path):
    from pptx import Presentation
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[5])
    slide.shapes.title.text = "FLEX DS Operating Model"
    p = tmp_path / "flex.pptx"
    pres.save(p)
    text, used = extract_text(p)
    assert "FLEX" in text
    assert used == "pptx"


def test_extract_text_pdf(tmp_path: Path):
    # Construct a tiny single-page PDF via pypdf
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    p = tmp_path / "empty.pdf"
    with open(p, "wb") as f:
        writer.write(f)
    text, used = extract_text(p)
    # Blank PDF — text is empty but extractor should succeed
    assert used == "pdf"
    assert isinstance(text, str)


def test_extract_text_missing_file(tmp_path: Path):
    text, used = extract_text(tmp_path / "nope.docx")
    assert text == ""
    assert used.startswith("skip:")


def test_extract_text_unknown_extension(tmp_path: Path):
    p = tmp_path / "weird.xyz"
    p.write_bytes(b"binary")
    text, used = extract_text(p)
    assert text == ""
    assert used.startswith("skip:")


def test_extract_text_corrupt_docx(tmp_path: Path):
    p = tmp_path / "corrupt.docx"
    p.write_bytes(b"this is not a zip file")
    text, used = extract_text(p)
    assert text == ""
    assert used.startswith("skip:")


# --- build_classification_prompt ---

def _rec_for_prompt(**overrides):
    base = {
        "index": 0,
        "role": "airbnb",
        "path": Path("E:/Dropbox/Archive/Example Corp/Roadmap.docx"),
        "size": 1000,
        "ext": ".docx",
        "text_excerpt": "Vision: improve host pricing.",
        "mtime": "2025-07-15T10:00:00+00:00",
    }
    base.update(overrides)
    return base


def test_build_classification_prompt_includes_all_classes():
    prompt = build_classification_prompt([_rec_for_prompt()])
    for cls in ALLOWED_CLASSES:
        assert cls in prompt
    assert "Roadmap.docx" in prompt
    assert "airbnb" in prompt
    assert "Vision: improve host pricing." in prompt


def test_build_classification_prompt_numbers_multiple_files():
    batch = [
        _rec_for_prompt(index=0, path=Path("A.docx"), text_excerpt="first"),
        _rec_for_prompt(index=1, path=Path("B.docx"), text_excerpt="second"),
    ]
    prompt = build_classification_prompt(batch)
    assert "[0]" in prompt
    assert "[1]" in prompt
    assert "first" in prompt
    assert "second" in prompt


def test_build_classification_prompt_instructs_json_object_output():
    prompt = build_classification_prompt([_rec_for_prompt()])
    # The parser expects {"classifications": [...]}
    assert "classifications" in prompt


def test_build_classification_prompt_shows_date_hints():
    """Classifier must see filename + mtime so it can infer chronological_date."""
    prompt = build_classification_prompt([_rec_for_prompt(
        path=Path("H2 2025 Marketplace DS Roadmap.docx"),
        mtime="2025-07-15T10:00:00+00:00",
    )])
    assert "H2 2025" in prompt
    assert "2025" in prompt   # mtime year at minimum
    # Prompt should request chronological_date in output
    assert "chronological_date" in prompt


def test_build_classification_prompt_uses_parent_date_hint_if_provided():
    """Workflowy walker passes parent_date (from date-heading node).
    The classifier should receive it as a high-confidence hint."""
    prompt = build_classification_prompt([_rec_for_prompt(
        path="wf://abc-123",
        parent_date="2019-06-04",
    )])
    assert "2019-06-04" in prompt


# --- parse_classification_response ---

def test_parse_classification_response_valid_json():
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "roadmap", "signal_score": 0.9, "summary": "H2 2025 plan",
             "chronological_date": "2025-07", "date_source": "filename", "date_confidence": 0.95},
            {"index": 1, "class": "bio", "signal_score": 0.7, "summary": "LinkedIn bio",
             "chronological_date": "2022-11", "date_source": "mtime", "date_confidence": 0.8},
        ]
    })
    result = parse_classification_response(raw)
    assert len(result) == 2
    assert result[0]["index"] == 0
    assert result[0]["class"] == "roadmap"
    assert result[0]["signal_score"] == 0.9
    assert result[0]["chronological_date"] == "2025-07"
    assert result[0]["date_source"] == "filename"
    assert result[0]["date_confidence"] == 0.95
    assert result[1]["class"] == "bio"
    assert result[1]["chronological_date"] == "2022-11"


def test_parse_classification_response_defaults_missing_date_fields():
    """Classifier may omit date fields on low-confidence records; parser fills defaults."""
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "ephemeral", "signal_score": 0.1, "summary": "scrap"},
        ]
    })
    result = parse_classification_response(raw)
    assert result[0]["chronological_date"] == ""
    assert result[0]["date_source"] == "unknown"
    assert result[0]["date_confidence"] == 0.0


def test_parse_classification_response_clamps_date_confidence():
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "bio", "signal_score": 0.5, "summary": "x",
             "chronological_date": "2020", "date_source": "content", "date_confidence": 1.7},
        ]
    })
    result = parse_classification_response(raw)
    assert result[0]["date_confidence"] == 1.0


def test_parse_classification_response_coerces_bad_date_source():
    """If LLM hallucinates a date_source, default to 'unknown'."""
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "bio", "signal_score": 0.5, "summary": "x",
             "chronological_date": "2020", "date_source": "stars", "date_confidence": 0.5},
        ]
    })
    result = parse_classification_response(raw)
    assert result[0]["date_source"] == "unknown"


def test_parse_classification_response_clamps_unknown_class_to_skip():
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "hallucinated_class", "signal_score": 0.5, "summary": "?"},
        ]
    })
    result = parse_classification_response(raw)
    assert result[0]["class"] == "skip"


def test_parse_classification_response_clamps_signal_score():
    raw = json.dumps({
        "classifications": [
            {"index": 0, "class": "bio", "signal_score": 2.5, "summary": "x"},
            {"index": 1, "class": "bio", "signal_score": -0.2, "summary": "y"},
        ]
    })
    result = parse_classification_response(raw)
    assert result[0]["signal_score"] == 1.0
    assert result[1]["signal_score"] == 0.0


def test_parse_classification_response_invalid_json_raises():
    with pytest.raises(ValueError):
        parse_classification_response("not json at all")


def test_parse_classification_response_missing_classifications_key_raises():
    with pytest.raises(ValueError):
        parse_classification_response(json.dumps({"other": []}))


def test_parse_classification_response_tolerates_missing_summary():
    raw = json.dumps({
        "classifications": [{"index": 0, "class": "bio", "signal_score": 0.8}],
    })
    result = parse_classification_response(raw)
    assert result[0]["summary"] == ""


# --- merge_records_into_index ---

def _rec(path_str: str, mtime: str, klass: str = "strategic_doc", role: str = "airbnb"):
    return {
        "path": path_str,
        "role": role,
        "mtime": mtime,
        "size": 100,
        "ext": ".docx",
        "class": klass,
        "signal_score": 0.8,
        "summary": "test",
        "classified_at": "2026-04-21T00:00:00Z",
    }


def test_merge_records_into_index_inserts_new():
    index = {"files": {}}
    merged = merge_records_into_index(index, [_rec("a.docx", "2026-04-01T00:00:00Z")])
    assert "a.docx" in merged["files"]
    assert merged["files"]["a.docx"]["class"] == "strategic_doc"


def test_merge_records_into_index_preserves_unchanged_entries():
    existing = {"files": {"old.docx": _rec("old.docx", "2026-01-01T00:00:00Z")}}
    merged = merge_records_into_index(existing, [_rec("new.docx", "2026-04-01T00:00:00Z")])
    assert "old.docx" in merged["files"]
    assert "new.docx" in merged["files"]


def test_merge_records_into_index_updates_when_mtime_changes():
    existing = {"files": {"a.docx": _rec("a.docx", "2026-01-01T00:00:00Z", klass="bio")}}
    new = [_rec("a.docx", "2026-04-01T00:00:00Z", klass="roadmap")]
    merged = merge_records_into_index(existing, new)
    assert merged["files"]["a.docx"]["class"] == "roadmap"
    assert merged["files"]["a.docx"]["mtime"] == "2026-04-01T00:00:00Z"


def test_merge_records_into_index_is_idempotent_on_same_mtime():
    existing = {"files": {"a.docx": _rec("a.docx", "2026-01-01T00:00:00Z", klass="bio")}}
    # Same mtime — should preserve existing classification, not overwrite
    new = [_rec("a.docx", "2026-01-01T00:00:00Z", klass="roadmap")]
    merged = merge_records_into_index(existing, new)
    assert merged["files"]["a.docx"]["class"] == "bio"


def test_merge_records_into_index_overwrite_bypasses_idempotency():
    """overwrite=True is used when re-classifying specific buckets after
    a prompt tightening — same-mtime records must be overwritten."""
    existing = {"files": {"a.docx": _rec("a.docx", "2026-01-01T00:00:00Z", klass="feedback_received")}}
    new = [_rec("a.docx", "2026-01-01T00:00:00Z", klass="feedback_given")]
    merged = merge_records_into_index(existing, new, overwrite=True)
    assert merged["files"]["a.docx"]["class"] == "feedback_given"


def test_merge_records_into_index_tracks_generated_at():
    merged = merge_records_into_index({"files": {}}, [_rec("a.docx", "2026-01-01T00:00:00Z")])
    assert "generated_at" in merged
    assert "T" in merged["generated_at"]


# --- Extension lists sanity ---

def test_allowed_extensions_includes_career_doc_types():
    for ext in [".txt", ".md", ".docx", ".pdf", ".xlsx", ".pptx", ".csv"]:
        assert ext in ALLOWED_EXTENSIONS


def test_skip_extensions_includes_common_media():
    for ext in [".jpg", ".jpeg", ".png", ".gif", ".mov", ".mp4", ".mp3"]:
        assert ext in SKIP_EXTENSIONS
